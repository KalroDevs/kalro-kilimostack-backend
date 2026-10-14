from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError, PdfReadError

from .ollama_client import ollama_client
from .text_processing import split_long_text
from .vector_store import vector_store

MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".json",
    ".txt",
    ".md",
    ".markdown",
}


def extract_document_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_DOCUMENT_EXTENSIONS))
        raise ValueError(f"Unsupported file type. Supported extensions: {supported}.")

    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            document = WordDocument(BytesIO(content))
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts)
        elif extension == ".xlsx":
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            parts = []
            try:
                for sheet in workbook.worksheets:
                    parts.append(f"Sheet: {sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        values = [str(value).strip() for value in row if value is not None]
                        if values:
                            parts.append(" | ".join(values))
            finally:
                workbook.close()
            text = "\n".join(parts)
        elif extension == ".json":
            parsed = json.loads(content.decode("utf-8-sig"))
            text = json.dumps(parsed, ensure_ascii=False, indent=2)
        else:
            text = content.decode("utf-8-sig")
    except (
        BadZipFile,
        FileNotDecryptedError,
        InvalidFileException,
        json.JSONDecodeError,
        PackageNotFoundError,
        PdfReadError,
        UnicodeDecodeError,
    ) as exc:
        raise ValueError(f"Could not parse {filename}: {exc}") from exc

    if not text.strip():
        raise ValueError(
            "No extractable text was found. Scanned/image-only PDFs require OCR, "
            "which is not enabled."
        )
    return text


async def ingest_document(
    document_id: str,
    name: str,
    filename: str,
    content: bytes,
) -> int:
    text = extract_document_text(filename, content)
    chunks = split_long_text(text)
    ids = [
        f"training-document-{document_id}-{index}"
        for index in range(len(chunks))
    ]
    embeddings = []
    metadatas = []
    for index, chunk in enumerate(chunks):
        embeddings.append(await ollama_client.embed(f"{name}\n\n{chunk}"))
        metadatas.append(
            {
                "document_id": document_id,
                "title": name,
                "file_name": Path(filename).name,
                "content_id": ids[index],
                "content_header": f"Chunk {index + 1}",
                "chunk_index": index,
            }
        )

    vector_store.delete_training_document(document_id)
    vector_store.upsert_document_chunks(
        ids=ids,
        embeddings=embeddings,
        documents=chunks,
        metadatas=metadatas,
    )
    return len(chunks)
