from io import BytesIO

import pytest
from docx import Document as WordDocument
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from pypdf import PdfWriter

from app import document_ingestion
from app.config import settings
from app.main import app
from app.ollama_client import OllamaClient
from app.vector_store import VectorStore


def test_extracts_plain_text_and_markdown():
    assert document_ingestion.extract_document_text("notes.txt", b"Crop care") == "Crop care"
    assert document_ingestion.extract_document_text("notes.md", b"# Crop care") == "# Crop care"


def test_extracts_json_as_readable_text():
    text = document_ingestion.extract_document_text("resource.json", b'{"crop":"maize"}')
    assert '"crop": "maize"' in text


def test_extracts_word_paragraphs_and_tables():
    document = WordDocument()
    document.add_paragraph("Crop care guide")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Stage"
    table.cell(0, 1).text = "Action"
    buffer = BytesIO()
    document.save(buffer)

    text = document_ingestion.extract_document_text("guide.docx", buffer.getvalue())
    assert "Crop care guide" in text
    assert "Stage | Action" in text


def test_extracts_spreadsheet_sheets_and_rows():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Potato"
    sheet.append(["Pest", "Control"])
    sheet.append(["Aphids", "Monitor the crop"])
    buffer = BytesIO()
    workbook.save(buffer)

    text = document_ingestion.extract_document_text("guide.xlsx", buffer.getvalue())
    assert "Sheet: Potato" in text
    assert "Aphids | Monitor the crop" in text


def test_rejects_scanned_pdf_without_extractable_text():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(ValueError, match="OCR"):
        document_ingestion.extract_document_text("scan.pdf", buffer.getvalue())


@pytest.mark.asyncio
async def test_document_ingestion_endpoint_accepts_multipart_upload(monkeypatch):
    async def fake_embed(self, text, model=None):
        return [0.1, 0.2]

    monkeypatch.setattr(OllamaClient, "embed", fake_embed)
    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://test",
    ) as client:
        response = await client.post(
            "/api/v1/documents/ingest",
            data={"document_id": "http-test", "name": "HTTP Test"},
            files={"file": ("http-test.txt", b"Text from an uploaded document.", "text/plain")},
        )

    assert response.status_code == 200
    assert response.json() == {
        "document_id": "http-test",
        "status": "verified",
        "chunks_indexed": 1,
    }


@pytest.mark.asyncio
async def test_ingest_writes_only_to_the_training_document_collection(
    monkeypatch, tmp_path
):
    async def fake_embed(self, text, model=None):
        return [0.1, 0.2]

    monkeypatch.setattr(settings, "chroma_persist_dir", str(tmp_path))
    training_store = VectorStore()
    monkeypatch.setattr(document_ingestion, "vector_store", training_store)
    monkeypatch.setattr(OllamaClient, "embed", fake_embed)

    chunks_indexed = await document_ingestion.ingest_document(
        document_id="42",
        name="Crop Notes",
        filename="crop-notes.txt",
        content=b"Crop notes for staff training.",
    )

    assert chunks_indexed == 1
    assert training_store._documents_collection.count() == 1
    assert training_store._collection.count() == 0
    assert training_store.query([0.1, 0.2], top_k=1)["ids"] == [[]]
