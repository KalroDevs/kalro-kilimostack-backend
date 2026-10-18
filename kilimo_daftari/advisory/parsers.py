# advisory/parsers.py
"""
Parse the Screen & Classify workbook into dictionaries suitable for
AdvisoryResource and ContentSection.
"""
import re
from datetime import datetime, date

import openpyxl

from .models import (
    ContentTypeChoices,
    SectorChoices,
    ValidationStatusChoices,
    CurrencyStatusChoices,
    AccuracyCheckChoices,
    RiskLevelChoices,
    QualityFlagChoices,
)


MULTI_SPLIT = re.compile(r"\s*;\s*|\s*\n\s*")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def clean(v):
    return "" if v is None else str(v).strip()


def split_list(v):
    s = clean(v)
    return [p for p in MULTI_SPLIT.split(s) if p] if s else []


def split_tags(v):
    s = clean(v)
    return [t.strip() for t in s.split(",") if t.strip()] if s else []


def parse_date(v):
    if not v:
        return None
    if isinstance(v, (datetime, date)):
        return v if isinstance(v, date) else v.date()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def enum_value(raw, choices_cls, aliases=None):
    s = clean(raw)
    if not s:
        return ""
    key = s.lower().replace("-", "_").replace(" ", "_")
    for value, _label in choices_cls.choices:
        if value == key:
            return value
    for value, label in choices_cls.choices:
        if label.lower() == s.lower():
            return value
    if aliases:
        aliased = aliases.get(key) or aliases.get(s.lower())
        if aliased:
            return aliased
    return ""


CONTENT_TYPE_ALIASES = {
    "production_guide":    ContentTypeChoices.FARMER_GUIDE,
    "protocol":            ContentTypeChoices.TECHNICAL_MANUAL,
    "protocol/sop":        ContentTypeChoices.TECHNICAL_MANUAL,
    "sop":                 ContentTypeChoices.TECHNICAL_MANUAL,
    "advisory_bulletin":   ContentTypeChoices.FACTSHEET,
    "brochure":            ContentTypeChoices.FACTSHEET,
    "poster":              ContentTypeChoices.FACTSHEET,
    "poster/infographic":  ContentTypeChoices.FACTSHEET,
    "infographic":         ContentTypeChoices.FACTSHEET,
    "video":               ContentTypeChoices.HTML,
}

SECTOR_FROM_GROUP = {
    "crop":      SectorChoices.CROPS,
    "livestock": SectorChoices.LIVESTOCK,
}


# ---------------------------------------------------------------------------
# Header discovery
# ---------------------------------------------------------------------------
def _header_row_inventory(ws):
    for r in range(1, 15):
        if clean(ws.cell(row=r, column=2).value) == "ID":
            return r
    return None


def _header_row_section(ws):
    for r in range(1, 20):
        for c in range(1, ws.max_column + 1):
            if clean(ws.cell(row=r, column=c).value).lower() == "section header":
                return r
    return None


# ---------------------------------------------------------------------------
# Parsers
# ---------------------------------------------------------------------------
def parse_inventory_sheet(ws, group):
    hrow = _header_row_inventory(ws)
    if hrow is None:
        return []

    sector_value = SECTOR_FROM_GROUP.get(group, "")
    out = []
    for r in range(hrow + 1, ws.max_row + 1):
        resource_id = clean(ws.cell(row=r, column=2).value)
        title       = clean(ws.cell(row=r, column=3).value)
        if not resource_id or not title:
            continue
        if resource_id.endswith("-EX"):
            continue

        source_link = clean(ws.cell(row=r, column=4).value)
        is_url = source_link.lower().startswith(("http://", "https://"))

        row = {
            "resource_id":       resource_id,
            "group":             group,
            "title":             title,
            "link":              source_link if is_url else "",
            "publication_id":    "" if is_url else source_link,
            "institution":       clean(ws.cell(row=r, column=5).value)
                                 or "Kenya Agricultural and Livestock Research Organization",
            "publish_date":      clean(ws.cell(row=r, column=6).value),
            "value_chain":       clean(ws.cell(row=r, column=7).value),
            "advisory_domain":   split_list(ws.cell(row=r, column=8).value),
            "content_type":      enum_value(
                                     ws.cell(row=r, column=9).value,
                                     ContentTypeChoices,
                                     CONTENT_TYPE_ALIASES,
                                 ),
            "production_system": split_list(ws.cell(row=r, column=11).value),
            "target_users":      split_list(ws.cell(row=r, column=12).value),
            "currency_status":   enum_value(ws.cell(row=r, column=15).value, CurrencyStatusChoices),
            "scientific_accuracy_check": enum_value(
                                             ws.cell(row=r, column=16).value, AccuracyCheckChoices
                                         ),
            "validation_status": enum_value(
                                     ws.cell(row=r, column=17).value, ValidationStatusChoices
                                 ),
            "risk_level":        enum_value(ws.cell(row=r, column=18).value, RiskLevelChoices),
            "quality_flag":      enum_value(ws.cell(row=r, column=19).value, QualityFlagChoices),
            "screening_notes":   clean(ws.cell(row=r, column=20).value),
        }
        if sector_value:
            row["sector"] = sector_value
        out.append(row)
    return out


def parse_section_sheet(ws, sheet_name):
    hrow = _header_row_section(ws)
    if hrow is None:
        return []

    headers = {
        clean(ws.cell(row=hrow, column=c).value).lower(): c
        for c in range(1, ws.max_column + 1)
        if clean(ws.cell(row=hrow, column=c).value)
    }
    c_num     = headers.get("#")
    c_header  = headers.get("section header")
    c_summary = headers.get("content summary")
    c_pages   = headers.get("pages")
    c_warn    = headers.get("content warnings")
    c_tags    = headers.get("content tags")

    if not (c_num and c_header and c_summary and c_pages):
        return []

    out = []
    for r in range(hrow + 1, ws.max_row + 1):
        num = ws.cell(row=r, column=c_num).value
        if num is None or clean(num) == "":
            continue
        try:
            order = int(float(num))
        except (TypeError, ValueError):
            continue

        header  = clean(ws.cell(row=r, column=c_header).value)
        summary = clean(ws.cell(row=r, column=c_summary).value)
        pages   = clean(ws.cell(row=r, column=c_pages).value)
        warn    = clean(ws.cell(row=r, column=c_warn).value) if c_warn else ""
        tags    = clean(ws.cell(row=r, column=c_tags).value) if c_tags else ""

        if not (header or summary) or not pages:
            continue

        out.append({
            "reading_order":    order,
            "content_header":   header,
            "content_text":     summary,
            "pages_reference":  pages,
            "value_chain":      sheet_name,
            "content_warnings": split_list(warn),
            "content_tags":     split_tags(tags),
        })
    return out
