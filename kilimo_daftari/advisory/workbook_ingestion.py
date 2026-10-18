"""Import advisory inventory and value chain sections from uploaded XLSX files."""
from __future__ import annotations

from collections import Counter
import re
from datetime import date, datetime

from django.db import transaction

from .models import (
    AdvisoryResource, ContentSection, ContentTypeChoices,
    SectorChoices, ValueChain, WorkbookUpload,
)


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header(value):
    return re.sub(r"\s+", " ", _text(value).replace("\n", " ").replace("\r", " ")).strip().lower()


def _split(value, comma=False):
    text = _text(value)
    if not text:
        return []
    return [part.strip() for part in re.split(r"[,;\n]" if comma else r"[;\n]", text) if part.strip()]


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _sheet_rows(ws):
    """Find a recognizable header row within the first 25 rows."""
    if ws.max_row is None or ws.max_column is None:
        # Read-only worksheets can omit their dimensions in the XLSX file.
        ws.calculate_dimension(force=True)
    max_row = ws.max_row or 0
    max_column = ws.max_column or 0
    if max_row < 1 or max_column < 1:
        return None, None, []

    rows = ws.iter_rows(
        min_row=1,
        max_row=min(max_row, 25),
        max_col=max_column,
        values_only=True,
    )
    for row_no, cells in enumerate(rows, 1):
        headers = [_header(v) for v in cells]
        present = {h for h in headers if h}
        if {"id", "content title"} <= present:
            return "inventory", row_no, headers
        if {"section header", "content summary", "pages"} <= present:
            return "content", row_no, headers
    return None, None, []


def _columns(headers):
    return {h: i for i, h in enumerate(headers) if h}


def _find(columns, *needles):
    for needle in needles:
        for name, index in columns.items():
            if needle in name:
                return index
    return None


def _cell(row, index):
    return row[index] if index is not None and index < len(row) else None


def _enum(value, choices, aliases=None, default=""):
    raw = _text(value).lower().replace("-", "_").replace(" ", "_")
    if not raw:
        return default
    valid = {choice[0] for choice in choices}
    if raw in valid:
        return raw
    for key, label in choices:
        if _text(label).lower() == _text(value).lower():
            return key
    return (aliases or {}).get(raw, default)


def _value_chain(name):
    existing = ValueChain.objects.filter(name__iexact=name).first()
    return existing or ValueChain.objects.create(name=name)


def _inventory(ws, header_row, headers):
    from providers.models import Provider
    cols = _columns(headers)
    required = {"id": _find(cols, "id"), "title": _find(cols, "content title")}
    if any(i is None for i in required.values()):
        raise ValueError(f"Worksheet '{ws.title}' needs ID and Content Title columns.")
    result = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rid, title = _text(_cell(values, required["id"])), _text(_cell(values, required["title"]))
        if not rid or not title:
            continue
        source = _text(_cell(values, _find(cols, "publication id", "source link")))
        institution_name = _text(_cell(values, _find(cols, "institution")))
        fmt = _text(_cell(values, _find(cols, "format"))).lower()
        content_type = _enum(_cell(values, _find(cols, "content type")), ContentTypeChoices.choices,
                             {"pdf": ContentTypeChoices.PDF, "html": ContentTypeChoices.HTML,
                              "production_guide": ContentTypeChoices.FARMER_GUIDE,
                              "protocol": ContentTypeChoices.TECHNICAL_MANUAL},
                             default=ContentTypeChoices.PDF)
        aez = _text(_cell(values, _find(cols, "country", "aez")))
        geo = {}
        if aez:
            country, sep, zone = aez.partition("-")
            geo = {"country": country.strip(), "agro_ecological_zones": [zone.strip()] if sep and zone.strip() else [], "notes": aez}
        provider = None
        if institution_name:
            provider = Provider.objects.filter(name__iexact=institution_name).first()
        sheet_key = _header(ws.title)
        sector = SectorChoices.LIVESTOCK if "livestock" in sheet_key else SectorChoices.CROPS if "crop" in sheet_key else SectorChoices.CROSS_CUTTING
        result.append({
            "workbook_resource_id": rid,
            "workbook_source_reference": source,
            "title": title,
            "link": source if source.lower().startswith(("http://", "https://")) else "",
            "institution": provider,
            "author": _split(_cell(values, _find(cols, "author")), comma=True),
            "publish_date": _text(_cell(values, _find(cols, "publish date"))),
            "content_type": content_type,
            "sector": sector,
            "advisory_domain": _split(_cell(values, _find(cols, "advisory domain", "topic"))),
            "production_system": _split(_cell(values, _find(cols, "production system"))),
            "target_users": _split(_cell(values, _find(cols, "target users"))),
            "geographic_applicability": geo,
            "review_date": _date(_cell(values, _find(cols, "last reviewed", "currency check"))),
            "currency_status": _enum(_cell(values, _find(cols, "currency status")), AdvisoryResource._meta.get_field("currency_status").choices),
            "scientific_accuracy_check": _enum(_cell(values, _find(cols, "scientific accuracy")), AdvisoryResource._meta.get_field("scientific_accuracy_check").choices),
            "validation_status": _enum(_cell(values, _find(cols, "validation status")), AdvisoryResource._meta.get_field("validation_status").choices),
            "risk_level": _enum(_cell(values, _find(cols, "risk level")), AdvisoryResource._meta.get_field("risk_level").choices),
            "quality_flag": _enum(_cell(values, _find(cols, "quality flag")), AdvisoryResource._meta.get_field("quality_flag").choices),
            "screening_notes": _text(_cell(values, _find(cols, "screening notes", "escalation"))),
            "_value_chain": _text(_cell(values, _find(cols, "value chain", "commodity"))),
        })
    return result


def _content(ws, header_row, headers):
    cols = _columns(headers)
    result = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        header = _text(_cell(values, _find(cols, "section header")))
        summary = _text(_cell(values, _find(cols, "content summary")))
        pages = _text(_cell(values, _find(cols, "pages")))
        if not (header or summary or pages):
            continue
        raw_order = _cell(values, _find(cols, "#"))
        try:
            order = int(float(raw_order)) if raw_order not in (None, "") else len(result) + 1
        except (ValueError, TypeError):
            order = len(result) + 1
        result.append({
            "reading_order": order, "content_header": header, "content_text": summary,
            "pages": pages,
            "content_warnings": _split(_cell(values, _find(cols, "content warnings"))),
            "content_tags": _split(_cell(values, _find(cols, "content tags")), comma=True),
            "status": _text(_cell(values, _find(cols, "status"))),
            "value_chain": ws.title,
        })
    return result


def _sync_imported_resources(upload_id, resource_ids):
    """Push validated workbook sections to the AI index after DB commit."""
    from .services import sync_validated_workbook_resources

    resources = AdvisoryResource.objects.filter(pk__in=resource_ids).prefetch_related(
        "content_sections", "value_chains"
    )
    result = sync_validated_workbook_resources(resources)
    upload = WorkbookUpload.objects.filter(pk=upload_id).first()
    if upload:
        summary = result.summary()
        upload.log = (
            f"{upload.log}\nAI index: {summary['synced']} synced, "
            f"{summary['skipped']} skipped, {summary['failed']} failed."
        ).strip()
        upload.save(update_fields=["log"])


@transaction.atomic
def _ingest_workbook(upload: WorkbookUpload):
    """Validate then atomically upsert inventory and linked content sections."""
    upload.status = WorkbookUpload.Status.PROCESSING
    upload.save(update_fields=["status"])
    counts = {"resources_created": 0, "resources_updated": 0, "sections_created": 0,
              "sections_updated": 0, "sections_skipped": 0}
    messages = []
    try:
        import openpyxl

        upload.file.open("rb")
        workbook = openpyxl.load_workbook(upload.file, data_only=True, read_only=True)
        inventories, content_sheets = [], []
        for ws in workbook.worksheets:
            role, header_row, headers = _sheet_rows(ws)
            if role == "inventory":
                rows = _inventory(ws, header_row, headers)
                inventories.extend(rows)
                messages.append(f"{ws.title}: {len(rows)} advisory resources")
            elif role == "content":
                rows = _content(ws, header_row, headers)
                content_sheets.extend((row, ws.title) for row in rows)
                messages.append(f"{ws.title}: {len(rows)} content rows")
        if not inventories:
            raise ValueError("No advisory resource worksheet found. Add a sheet with ID and Content Title headers.")
        ids = [row["workbook_resource_id"] for row in inventories]
        duplicates = sorted(resource_id for resource_id, count in Counter(ids).items() if count > 1)
        if duplicates:
            raise ValueError(
                "The workbook contains duplicate advisory IDs: "
                f"{', '.join(duplicates)}. Each ID must be unique."
            )
        # Upsert resources first so section Pages references resolve independent of sheet order.
        resources_by_id = {}
        for item in inventories:
            data = {k: v for k, v in item.items() if not k.startswith("_")}
            obj, created = AdvisoryResource.objects.update_or_create(
                workbook_resource_id=item["workbook_resource_id"], defaults=data
            )
            resources_by_id[item["workbook_resource_id"]] = obj
            counts["resources_created" if created else "resources_updated"] += 1
            names = [n for n in (item["_value_chain"],) if n]
            for name in names:
                vc = _value_chain(name)
                obj.value_chains.add(vc)
        for item, sheet_name in content_sheets:
            parent = resources_by_id.get(item["pages"])
            if parent is None:
                parent = AdvisoryResource.objects.filter(workbook_resource_id=item["pages"]).first()
            if parent is None:
                counts["sections_skipped"] += 1
                messages.append(f"{sheet_name}: Pages '{item['pages']}' did not match an advisory ID")
                continue
            vc = _value_chain(sheet_name)
            parent.value_chains.add(vc)
            content_id = f"{item['pages']}-S{item['reading_order']:04d}"
            defaults = {"resource": parent, "reading_order": item["reading_order"],
                        "pages_reference": item["pages"],
                        "content_header": item["content_header"], "content_text": item["content_text"],
                        "content_warnings": item["content_warnings"], "content_tags": item["content_tags"]}
            if item["status"]:
                defaults["status"] = (
                    _enum(item["status"], ContentSection._meta.get_field("status").choices)
                    or item["status"][:30]
                )
            section = ContentSection.objects.filter(
                resource=parent, reading_order=item["reading_order"]
            ).first()
            if section is None:
                section, created = ContentSection.objects.update_or_create(
                    content_id=content_id, defaults=defaults
                )
            else:
                section.content_id = content_id
                for field, value in defaults.items():
                    setattr(section, field, value)
                section.save()
                created = False
            section.value_chains.add(vc)
            counts["sections_created" if created else "sections_updated"] += 1
        workbook.close()
        upload.status = WorkbookUpload.Status.PROCESSED
        upload.log = "\n".join(messages + [f"Unmatched content rows skipped: {counts['sections_skipped']}."])
    except Exception as exc:
        upload.status = WorkbookUpload.Status.FAILED
        upload.log = f"Import failed: {exc}"
        raise
    finally:
        upload.file.close()
    for key, value in counts.items():
        setattr(upload, key, value)
    upload.save(update_fields=["status", "log", *counts.keys()])
    resource_ids = tuple(resource.pk for resource in resources_by_id.values())
    transaction.on_commit(lambda: _sync_imported_resources(upload.pk, resource_ids))
    return counts


def ingest_workbook(upload: WorkbookUpload):
    try:
        return _ingest_workbook(upload)
    except Exception as exc:
        upload.status = WorkbookUpload.Status.FAILED
        upload.log = f"Import failed: {exc}"
        upload.save(update_fields=["status", "log"])
        raise
