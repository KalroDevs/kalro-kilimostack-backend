import json
from unittest.mock import patch
from io import BytesIO
from pathlib import Path

from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
import httpx

from .models import AdvisoryResource, ContentSection, ValueChain, WorkbookUpload
from .serializers import AdvisoryResourceSerializer

SAMPLE_JSON = Path(__file__).resolve().parent.parent.parent / "data" / "sample_camel_calf_resource.json"


class AdvisoryResourceImportTests(TestCase):
    """
    Round-trips the real KALRO camel-calf example (Advisory Content Import
    JSON Specification v0.1) through the serializer, matching the manual
    verification used to develop this project.
    """

    def setUp(self):
        self.data = json.loads(SAMPLE_JSON.read_text())[0]

    def test_import_creates_resource_with_all_content_sections(self):
        serializer = AdvisoryResourceSerializer(data=self.data)
        self.assertTrue(serializer.is_valid(), serializer.errors)
        resource = serializer.save()

        self.assertEqual(resource.publication_id, "kalro-livestock-camel-calf-management-2017-043")
        self.assertEqual(resource.sector, "livestock")
        self.assertEqual(resource.value_chain, "camel")
        self.assertEqual(resource.content_sections.count(), 9)
        self.assertEqual(resource.risk_level, "high")
        self.assertTrue(resource.requires_human_review)

    def test_serialized_output_round_trips_spec_shape(self):
        serializer = AdvisoryResourceSerializer(data=self.data)
        serializer.is_valid(raise_exception=True)
        resource = serializer.save()

        output = AdvisoryResourceSerializer(resource).data
        self.assertEqual(len(output["content"]), 9)
        self.assertEqual(output["advisory_safety"]["risk_level"], "high")
        self.assertIn("arid and semi-arid lands", output["geographic_applicability"]["agro_ecological_zones"])

    def test_default_quality_flag_excludes_resource_from_ai_layer_sync(self):
        serializer = AdvisoryResourceSerializer(data=self.data)
        serializer.is_valid(raise_exception=True)
        resource = serializer.save()
        # Freshly imported resources haven't been screened yet.
        self.assertFalse(resource.is_ready_for_ai_layer)

        resource.quality_flag = "ready_to_certify"
        resource.save()
        self.assertTrue(resource.is_ready_for_ai_layer)


class ScreeningWorkflowTests(TestCase):
    def test_upsert_by_publication_id(self):
        data = json.loads(SAMPLE_JSON.read_text())[0]
        AdvisoryResourceSerializer(data=data).is_valid(raise_exception=True)
        first = AdvisoryResourceSerializer(data=data)
        first.is_valid(raise_exception=True)
        first.save()

        self.assertEqual(AdvisoryResource.objects.count(), 1)
        existing = AdvisoryResource.objects.get(publication_id=data["publication_id"])
        second = AdvisoryResourceSerializer(existing, data=data, partial=True)
        second.is_valid(raise_exception=True)
        second.save()

        self.assertEqual(AdvisoryResource.objects.count(), 1)


class WorkbookUploadImportTests(TestCase):
    def setUp(self):
        override = override_settings(STORAGES={
            "default": {"BACKEND": "django.core.files.storage.InMemoryStorage"},
            "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
        })
        override.enable()
        self.addCleanup(override.disable)

    def _workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Crop Resources"
        ws.append(["ID", "Content Title", "Publication ID / Source Link", "Institution & Author(s)",
                   "Publish Date", "Value Chain / Commodity", "Advisory Domain / Topic", "Content Type",
                   "Format", "Production System", "Target Users", "Country / AEZ", "Last Reviewed (Currency Check)",
                   "Currency Status", "Scientific Accuracy Check", "Validation Status", "Risk Level",
                   "Quality Flag", "Screening Notes / Escalation Guidance"])
        ws.append([101, "Bean production", "PUB-101", "KALRO; Dr Author", "2024-01-02", "Beans",
                   "Plant health", "Factsheet", "PDF", "Rainfed", "Farmers", "Kenya - Central",
                   "2025-01-01", "Current", "Verified", "Expert Reviewed", "Low", "Needs Review", "Check soil"])
        ws.append([None] * 19)
        livestock = wb.create_sheet("Livestock resources")
        livestock.append(["ID", "Content Title", "Publication ID / Source Link"])
        livestock.append(["L-202", "Dairy health", "https://example.org/dairy"])
        sections = wb.create_sheet("Novel Value Chain Name")
        sections.append(["#", "Section Header", "Content Summary", "Pages", "Content Warnings", "Content Tags", "Status"])
        sections.append([1, "Overview", "Beans summary", 101, "Caution", "beans, production", "Validated"])
        sections.append([2, "Unmatched", "No resource", 999, None, None, None])
        sections.append([None] * 7)
        out = BytesIO()
        wb.save(out)
        return out.getvalue()

    def _upload(self):
        return WorkbookUpload.objects.create(file=SimpleUploadedFile(
            "advisory.xlsx", self._workbook(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ))

    def test_imports_dynamic_sheets_and_links_pages_to_resource_id(self):
        upload = self._upload()
        self.assertEqual(upload.status, WorkbookUpload.Status.PROCESSED)
        self.assertEqual(upload.resources_created, 2)
        self.assertEqual(upload.sections_created, 1)
        self.assertEqual(upload.sections_skipped, 1)
        resource = AdvisoryResource.objects.get(workbook_resource_id="101")
        self.assertEqual(resource.workbook_source_reference, "PUB-101")
        section = ContentSection.objects.get(resource=resource)
        self.assertEqual(section.pages_reference, "101")
        self.assertEqual(section.content_id, "101-S0001")
        self.assertFalse(section.content_id.startswith("wb"))
        self.assertEqual(section.content_text, "Beans summary")
        self.assertEqual(list(section.value_chains.values_list("name", flat=True)), ["Novel Value Chain Name"])
        self.assertTrue(ValueChain.objects.filter(name="Beans").exists())
        self.assertIn("Pages '999'", upload.log)

    def test_reupload_updates_existing_records(self):
        self._upload()
        second = self._upload()
        self.assertEqual(second.status, WorkbookUpload.Status.PROCESSED)
        self.assertEqual(second.resources_created, 0)
        self.assertEqual(second.resources_updated, 2)
        self.assertEqual(second.sections_created, 0)
        self.assertEqual(second.sections_updated, 1)
        self.assertEqual(AdvisoryResource.objects.count(), 2)
        self.assertEqual(ContentSection.objects.count(), 1)

    def test_invalid_workbook_reports_actionable_failure(self):
        wb = Workbook()
        wb.active.append(["Unexpected", "Headers"])
        out = BytesIO()
        wb.save(out)
        upload = WorkbookUpload.objects.create(file=SimpleUploadedFile("invalid.xlsx", out.getvalue()))
        self.assertEqual(upload.status, WorkbookUpload.Status.FAILED)
        self.assertIn("No advisory resource worksheet", upload.log)


class AskEndpointTests(TestCase):
    @patch("advisory.views.httpx.post")
    def test_ask_proxies_to_ai_layer_chat(self, post):
        post.return_value = httpx.Response(
            200,
            json={"answer": "Use the approved guidance.", "sources": [], "model": "test"},
            request=httpx.Request("POST", "http://localhost:8001/api/v1/chat"),
        )
        response = self.client.post(
            "/api/v1/ask/",
            data=json.dumps({"query": "  How do I care for calves?  ", "filters": {"sector": "livestock"}}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["answer"], "Use the approved guidance.")
        args, kwargs = post.call_args
        self.assertTrue(args[0].endswith("/api/v1/chat"))
        self.assertEqual(kwargs["json"]["query"], "How do I care for calves?")
        self.assertEqual(kwargs["json"]["filters"], {"sector": "livestock"})

    def test_ask_rejects_missing_query(self):
        response = self.client.post("/api/v1/ask/", data={}, content_type="application/json")
        self.assertEqual(response.status_code, 400)


class WorkbookAiSyncTests(TestCase):
    @override_settings(
        AI_LAYER_SYNC_ENABLED=True,
        AI_LAYER_BASE_URL="http://ai-layer.test",
        AI_LAYER_SYNC_TIMEOUT_SECONDS=5,
        AI_LAYER_WORKBOOK_SYNC_TIMEOUT_SECONDS=300,
        AI_LAYER_SYNC_BATCH_SIZE=25,
    )
    @patch("advisory.services.httpx.post")
    def test_validated_workbook_sections_sync_with_ai_schema_and_source_ids(self, post):
        from .services import sync_validated_workbook_resources

        post.return_value = httpx.Response(
            200,
            json={"ingested": ["SOURCE-1"], "chunks_indexed": 1, "skipped": [], "errors": []},
            request=httpx.Request("POST", "http://ai-layer.test/api/v1/ingest"),
        )
        resource = AdvisoryResource.objects.create(
            title="Beans field guide",
            workbook_resource_id="BEAN-1",
            workbook_source_reference="SOURCE-1",
            content_type="farmer_guide",
            sector="crops",
        )
        value_chain = ValueChain.objects.create(name="Green Gram")
        resource.value_chains.add(value_chain)
        section = ContentSection.objects.create(
            resource=resource,
            content_id="BEAN-1-S0001",
            pages_reference="BEAN-1",
            reading_order=1,
            content_header="Seed selection",
            content_text="Select clean, viable seed before planting.",
            status="validated",
        )
        section.value_chains.add(value_chain)

        result = sync_validated_workbook_resources([resource])

        self.assertEqual(result.synced, 1)
        payload = post.call_args.kwargs["json"][0]
        self.assertEqual(payload["publication_id"], "SOURCE-1")
        self.assertEqual(payload["value_chain"], "Green Gram")
        self.assertEqual(payload["content"][0]["content_text"], section.content_text)
        self.assertEqual(resource.__class__.objects.get(pk=resource.pk).vector_sync_status, "synced")
